"""Build outputs/bibliography.xlsx from data/processed/records.json and data/raw/counts/.

Tabs: Read me · Records · Gap matrix · Growth · Emerging terms · Landmarks by theme ·
India subset · Shortlist. The Gap matrix uses COUNTIFS over whole Records columns, so it
updates when rows are added to Records.
"""
import datetime as dt
import json
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
recs = json.loads((ROOT / "data" / "processed" / "records.json").read_text())
cfg = yaml.safe_load((ROOT / "queries.yaml").read_text())
THEMES = cfg["themes"]
REGIONS = ["Global North", "China & East Asia", "South Asia", "Southeast Asia", "Africa",
           "Latin America", "Middle East", "Multi-region", "Not place-specific"]
METHODS = ["qualitative", "quantitative", "mixed", "review", "conceptual", "computational",
           "unclear"]
PERIODS = [("2010–14", 2010, 2014), ("2015–19", 2015, 2019), ("2020–26", 2020, 2026)]
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2F4F6F")
SUB = PatternFill("solid", fgColor="DCE6F0")
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def header(ws, row, values, fill=FILL, font=HEAD):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font = fill, font
        c.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


# ------------------------------------------------------------------ Records
cols = list(recs[0].keys())
ws_r = wb.active
ws_r.title = "Records"
header(ws_r, 1, cols)
for r in recs:
    ws_r.append([r[c] for c in cols])
ws_r.freeze_panes = "C2"
ws_r.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(recs) + 1}"
wmap = {"Title": 60, "Authors": 35, "Venue": 30, "One-line summary": 60, "Stated gaps": 60,
        "Places studied": 25, "Link": 30, "Open-access link": 30, "Found in": 30}
widths(ws_r, [wmap.get(c, 14) for c in cols])
L = {c: get_column_letter(i) for i, c in enumerate(cols, 1)}


def col(name):
    return f"Records!${L[name]}:${L[name]}"


# ------------------------------------------------------------------ Gap matrix
ws_g = wb.create_sheet("Gap matrix")
ws_g["A1"] = "Gap matrix — live COUNTIFS over the Records tab (counts update when rows are added)"
ws_g["A1"].font = Font(bold=True, size=12)
ws_g["A2"] = ("Stage 1 pilot: only C2, A9 and B5 have records; other themes show 0 until Stage 2. "
              "Counts are of records in this bibliography (a sample), not of all literature; "
              "use the Growth tab for OpenAlex-wide counts.")
row = 4


def block(title, colvals, crit):
    """crit(col_value) -> list of extra COUNTIFS criteria strings."""
    global row
    ws_g.cell(row=row, column=1, value=title).font = Font(bold=True)
    row += 1
    header(ws_g, row, ["Theme", "Name"] + colvals + ["Total", "India flag"], fill=SUB,
           font=Font(bold=True))
    hdr = row
    row += 1
    first = row
    for code, t in THEMES.items():
        ws_g.cell(row=row, column=1, value=code)
        ws_g.cell(row=row, column=2, value=t["name"])
        for j, v in enumerate(colvals, 3):
            cl = get_column_letter(j)
            extra = ",".join(crit(f"{cl}${hdr}"))
            ws_g.cell(row=row, column=j,
                      value=f"=COUNTIFS({col('Primary theme')},$A{row},{extra})")
        tc = get_column_letter(len(colvals) + 3)
        ws_g.cell(row=row, column=len(colvals) + 3,
                  value=f"=COUNTIFS({col('Primary theme')},$A{row})")
        ws_g.cell(row=row, column=len(colvals) + 4,
                  value=f"=COUNTIFS({col('Primary theme')},$A{row},{col('India flag')},\"Y\")")
        row += 1
    ws_g.cell(row=row, column=2, value="All themes").font = Font(bold=True)
    for j in range(3, len(colvals) + 5):
        cl = get_column_letter(j)
        ws_g.cell(row=row, column=j, value=f"=SUM({cl}{first}:{cl}{row - 1})").font = Font(bold=True)
    row += 3


block("Theme × region", REGIONS, lambda h: [f"{col('Region')},{h}"])
block("Theme × method (inferred from abstract)", METHODS, lambda h: [f"{col('Method')},{h}"])
# periods: header cell holds the label; bounds are written in the formula
ws_g.cell(row=row, column=1, value="Theme × period (publication year)").font = Font(bold=True)
row += 1
header(ws_g, row, ["Theme", "Name"] + [p[0] for p in PERIODS] + ["Before 2010", "Total"],
       fill=SUB, font=Font(bold=True))
row += 1
first = row
for code, t in THEMES.items():
    ws_g.cell(row=row, column=1, value=code)
    ws_g.cell(row=row, column=2, value=t["name"])
    for j, (_, a, b) in enumerate(PERIODS, 3):
        ws_g.cell(row=row, column=j, value=(
            f"=COUNTIFS({col('Primary theme')},$A{row},{col('Year')},\">={a}\","
            f"{col('Year')},\"<={b}\")"))
    ws_g.cell(row=row, column=6, value=(
        f"=COUNTIFS({col('Primary theme')},$A{row},{col('Year')},\"<2010\")"))
    ws_g.cell(row=row, column=7, value=f"=COUNTIFS({col('Primary theme')},$A{row})")
    row += 1
ws_g.cell(row=row, column=2, value="All themes").font = Font(bold=True)
for j in range(3, 8):
    cl = get_column_letter(j)
    ws_g.cell(row=row, column=j, value=f"=SUM({cl}{first}:{cl}{row - 1})").font = Font(bold=True)
widths(ws_g, [8, 48] + [13] * 11)
ws_g.freeze_panes = "C4"

# ------------------------------------------------------------------ Growth
ws_w = wb.create_sheet("Growth")
ws_w["A1"] = "Growth — OpenAlex-wide hit counts for each theme's combined query (not just this sample)"
ws_w["A1"].font = Font(bold=True, size=12)
ws_w["A2"] = ("Source: OpenAlex group_works by year, same query and mode as the search "
              "(data/raw/counts/). Ratio = mean per year 2020–26 ÷ mean per year 2015–19. "
              "Caveats: hit counts include off-topic works (precision of the full result set is "
              "well below that of the screened top results), and 2025–26 counts jump in every "
              "theme, partly because OpenAlex indexed many more repository/preprint records "
              "(e.g. Zenodo) in those years; Stage 3 will normalise by a baseline.")
ws_w["A2"].alignment = WRAP
ws_w.row_dimensions[2].height = 75
ws_w.merge_cells("A2:K2")
header(ws_w, 4, ["Theme", "Name", "2010–14", "2015–19", "2020–26", "Growth ratio (per-year)",
                 "SA 2010–14", "SA 2015–19", "SA 2020–26", "SA growth ratio",
                 "SA share 2020–26"])
r0 = 5
yearly = {}
for i, (code, t) in enumerate(THEMES.items()):
    rr = r0 + i
    ws_w.cell(row=rr, column=1, value=code)
    ws_w.cell(row=rr, column=2, value=t["name"])
    p = ROOT / "data" / "raw" / "counts" / f"{code}_by_year.json"
    if not p.exists():
        ws_w.cell(row=rr, column=3, value="Stage 2")
        continue
    d = json.loads(p.read_text())
    yearly[code] = d
    by = dict(zip(d["years"], d["global"]))
    sa = dict(zip(d["years"], d["south_asia"]))
    for j, (_, a, b) in enumerate(PERIODS, 3):
        ws_w.cell(row=rr, column=j, value=sum(by[y] for y in range(a, b + 1)))
        ws_w.cell(row=rr, column=j + 4, value=sum(sa[y] for y in range(a, b + 1)))
    ws_w.cell(row=rr, column=6, value=f"=IFERROR((E{rr}/7)/(D{rr}/5),\"\")").number_format = "0.00"
    ws_w.cell(row=rr, column=10, value=f"=IFERROR((I{rr}/7)/(H{rr}/5),\"\")").number_format = "0.00"
    ws_w.cell(row=rr, column=11, value=f"=IFERROR(I{rr}/E{rr},\"\")").number_format = "0.0%"
rr = r0 + len(THEMES) + 2
ws_w.cell(row=rr, column=1, value="Per-year counts (pilot themes)").font = Font(bold=True)
rr += 1
years = list(range(2010, 2027))
header(ws_w, rr, ["Theme", "Scope"] + years, fill=SUB, font=Font(bold=True))
for code, d in yearly.items():
    for scope in ("global", "south_asia"):
        rr += 1
        ws_w.cell(row=rr, column=1, value=code)
        ws_w.cell(row=rr, column=2, value=scope)
        for j, v in enumerate(d[scope], 3):
            ws_w.cell(row=rr, column=j, value=v)
widths(ws_w, [8, 44] + [11] * 17)

# ------------------------------------------------------------------ Emerging terms
ws_e = wb.create_sheet("Emerging terms")
ws_e["A1"] = "Emerging terms — Stage 3"
ws_e["A1"].font = Font(bold=True, size=12)
ws_e["A2"] = ("Top 30 keywords/topics whose frequency rose fastest from 2018–21 to 2022–26, with "
              "counts and example papers. To be filled in Stage 3.")
header(ws_e, 4, ["Rank", "Term", "Count 2018–21", "Count 2022–26", "Growth", "Themes",
                 "Example paper 1 (ID)", "Example paper 2 (ID)", "Example paper 3 (ID)"])
widths(ws_e, [6, 30, 14, 14, 10, 20, 22, 22, 22])

# ------------------------------------------------------------------ Landmarks, India, Shortlist
SHORT = ["ID", "Primary theme", "Title", "Year", "Authors", "Venue", "Cited-by count", "DOI",
         "Region", "Places studied", "Found in"]


def subset(name, rows, note, extra=()):
    ws = wb.create_sheet(name)
    ws["A1"] = note
    ws["A1"].alignment = WRAP
    ws.merge_cells(f"A1:{get_column_letter(len(SHORT) + len(extra))}1")
    ws.row_dimensions[1].height = 32
    header(ws, 2, SHORT + list(extra))
    for r in rows:
        ws.append([r[c] for c in SHORT + list(extra)])
    ws.freeze_panes = "D3"
    widths(ws, [10, 8, 60, 7, 30, 28, 9, 26, 16, 22, 26] + [30] * len(extra))
    return ws


subset("Landmarks by theme",
       sorted([r for r in recs if r["Landmark flag"] == "Y"],
              key=lambda r: (r["Primary theme"], -(r["Cited-by count"] or 0))),
       "Records found in a theme's 'landmarks' slice (top cited works for the query, any year) "
       "and kept after screening; sorted by theme, then citations.")
subset("India subset",
       sorted([r for r in recs if r["India flag"] == "Y"],
              key=lambda r: (r["Primary theme"], -(r["Year"] or 0))),
       "Records whose title/abstract name India or an Indian place (Delhi/NCR and Haryana flags "
       "in the extra columns).", extra=("Delhi/NCR flag", "Haryana flag"))
subset("Shortlist",
       [r for r in recs if r["Shortlist tag"]],
       "Records touching the researcher's current leads (stations, waiting, night-time "
       "mobility, fare integration/NCMC, rail-led urbanism, elevated infrastructure, Haryana "
       "secondary cities). Tagging only; it does not affect search or screening.",
       extra=("Shortlist tag",))

# ------------------------------------------------------------------ Read me
ws = wb.create_sheet("Read me", 0)
n = len(recs)
by_theme = {}
for r in recs:
    by_theme[r["Primary theme"]] = by_theme.get(r["Primary theme"], 0) + 1
lines = [
    ("Literature map: urban, digital and urban–digital research", "title"),
    (f"Stage 1 (pilot) build, {dt.date.today().isoformat()}. {n} deduplicated records "
     f"({', '.join(f'{k}: {v}' for k, v in sorted(by_theme.items()))}).", ""),
    ("What this is", "h"),
    ("A representative, reproducible map — not a census — of three domains (A urban, B digital "
     "society, C urban × digital), 48 themes. Every record comes from an OpenAlex API response; "
     "raw responses are cached in data/raw/ of the repository. See BRIEF.md, queries.yaml, "
     "PROGRESS.md and scripts/.", ""),
    ("Tabs", "h"),
    ("Records — one row per work (schema below). Gap matrix — live COUNTIFS of theme × region, "
     "theme × method, theme × period. Growth — OpenAlex-wide counts per theme and period. "
     "Emerging terms — Stage 3. Landmarks by theme — records from the 'landmarks' slice. "
     "India subset — India-flagged records. Shortlist — records touching current leads.", ""),
    ("How records were found", "h"),
    ("For each theme, 2–5 queries (queries.yaml) were OR-ed and run in OpenAlex exact-phrase "
     "mode over title, abstract and keywords, in three slices: landmarks (most cited, any year, "
     "top 50; page 2 added where fewer than 40 survived screening), recent (2022–2026, "
     "relevance-ranked, top 50), India/South Asia (same queries AND South Asian place names, "
     "relevance-ranked, top 50). Searches ran through the OpenAlex connector; full records "
     "were then fetched one by one from api.openalex.org by a single throttled fetcher.", ""),
    ("Screening", "h"),
    ("Rule-based on title + abstract: the theme's core concept must appear in the title or at "
     "least twice in the abstract, plus its context terms. One mention, a missing or stub "
     "abstract, or a technical/biomedical flag sends the record to manual review "
     "(data/screening/manual_review.csv, with a note per decision). Book reviews, errata, "
     "paratext, editorials, datasets and retracted works are excluded. Deduplicated by DOI, "
     "then by normalised title (the version with a publisher DOI and most citations is kept).",
     ""),
    ("Columns that are inferred automatically (check before citing)", "h"),
    ("Places studied / Region: place names matched in title + abstract (not author "
     "affiliations). Region buckets: Global North (Europe, North America, Australia, NZ, "
     "Russia), China & East Asia, South Asia, Southeast Asia (incl. Pacific islands), Africa "
     "(incl. North Africa), Latin America (incl. Caribbean), Middle East (incl. Central Asia "
     "and Caucasus); two or more buckets = Multi-region; none = Not place-specific (also used "
     "when there is no abstract). India / Delhi-NCR / Haryana flags: same matching; NCR "
     "includes Haryana NCR districts (e.g. Rohtak, Sonipat, Panipat). Method: keyword cues in "
     "the abstract; 'unclear' when there is no abstract or no cue. Landmark flag: found in the "
     "landmarks slice. Emerging flag (provisional): published 2022+ with field-weighted "
     "citation impact ≥ 1.5; Stage 3 replaces this with term-growth analysis. One-line summary "
     "and Stated gaps: sentences copied from the abstract, marked 'auto:'. Screening "
     "confidence: high = core concept in title and rule-included; medium = rule-included via "
     "abstract or kept on manual review; low = no abstract.", ""),
    ("Known limits of the pilot", "h"),
    ("Only C2, A9 and B5 have been run; the Gap matrix shows 0 elsewhere. A record's primary "
     "theme is chosen among the themes whose searches found it, so pilot records that belong "
     "to un-run themes (e.g. A12, A13, B1) sit in a pilot theme for now. Precision on a "
     "random 50 records: see data/screening/precision_sample_stage1.csv and PROGRESS.md.", ""),
    ("Extra columns", "h"),
    ("Found in = theme:slice combinations that returned the work; Duplicates merged = other "
     "OpenAlex IDs folded into this record; FWCI = OpenAlex field-weighted citation impact.",
     ""),
]
r = 1
for text, kind in lines:
    c = ws.cell(row=r, column=1, value=text)
    c.alignment = WRAP
    if kind == "title":
        c.font = Font(bold=True, size=14)
    elif kind == "h":
        c.font = Font(bold=True)
        c.fill = SUB
    else:
        ws.row_dimensions[r].height = max(15, 15 * (len(text) // 110 + 1))
    r += 1
ws.column_dimensions["A"].width = 130

out = ROOT / "outputs" / "bibliography.xlsx"
out.parent.mkdir(exist_ok=True)
wb.save(out)
print(out, n, "records")
